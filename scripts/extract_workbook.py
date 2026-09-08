#!/usr/bin/env python3
"""Extract a deterministic workbook snapshot without changing the workbook."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def extract(path: Path) -> dict:
    formulas = load_workbook(path, data_only=False, read_only=False)
    cached = load_workbook(path, data_only=True, read_only=False)
    sheets = []
    for ws in formulas.worksheets:
        cached_ws = cached[ws.title]
        cells = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                item = {"value": cell.value}
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    item["cached_value"] = cached_ws[cell.coordinate].value
                cells[cell.coordinate] = item
        sheets.append({"name": ws.title, "cells": cells})
    return {"workbook": path.name, "sha256": sha256(path), "tab_order": formulas.sheetnames, "sheets": sheets}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(extract(args.workbook), indent=2, sort_keys=True, default=str) + "\n", encoding="utf-8")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

