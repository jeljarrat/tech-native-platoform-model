#!/usr/bin/env python3
"""Build the 54-test Base package-repair release gate report."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import yaml
from openpyxl import load_workbook


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--qa-report", type=Path, required=True)
    parser.add_argument("--package-report", type=Path, required=True)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--spec", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    qa = json.loads(args.qa_report.read_text(encoding="utf-8"))
    package = json.loads(args.package_report.read_text(encoding="utf-8"))
    spec = yaml.safe_load(args.spec.read_text(encoding="utf-8"))
    workbook = load_workbook(args.workbook, data_only=True, read_only=False)
    tests = [{"test": f"qa::{item['test']}", "passed": item["passed"], "details": item["details"]} for item in qa["results"]]
    tests.extend({"test": f"package::{name}", "passed": passed, "details": name} for name, passed in package["checks"].items())
    for name, item in spec["golden_outputs"].items():
        actual = workbook[item["sheet"]][item["cell"]].value
        passed = abs(actual - item["value"]) <= item["absolute_tolerance"]
        tests.append({"test": f"golden::{name}", "passed": passed, "details": f"actual={actual}; expected={item['value']}"})
    for index, expected in enumerate(spec["tab_order"]):
        actual = workbook.sheetnames[index] if index < len(workbook.sheetnames) else None
        tests.append({"test": f"tab_order::{index + 1:02d}", "passed": actual == expected, "details": f"actual={actual}; expected={expected}"})
    if len(tests) != 54:
        raise RuntimeError(f"Release gate must contain exactly 54 tests; found {len(tests)}")
    payload = {"passed": all(item["passed"] for item in tests), "total_tests": len(tests), "failures": sum(not item["passed"] for item in tests), "tests": tests}
    args.output.write_text(json.dumps(payload, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(json.dumps({key: payload[key] for key in ("passed", "total_tests", "failures")}, indent=2))
    return 0 if payload["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
