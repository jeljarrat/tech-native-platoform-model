#!/usr/bin/env python3
"""Independently verify a shared-formula package repair."""

from __future__ import annotations

import argparse
import ast
import hashlib
import json
import math
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook


MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
ERRORS = {"#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!"}
LINEAGE_RANGES = {
    "Acquisition Schedule": ("O15:U15", "G40:P40"),
    "Org & Capacity": ("G116:N116",),
}


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def workbook_cells(path: Path, *, data_only: bool) -> dict[tuple[str, str], object]:
    workbook = load_workbook(path, data_only=data_only, read_only=False)
    return {(ws.title, cell.coordinate): cell.value for ws in workbook.worksheets for row in ws.iter_rows() for cell in row if cell.value is not None}


def formula_map(path: Path) -> dict[tuple[str, str], str]:
    workbook = load_workbook(path, data_only=False, read_only=False)
    return {(ws.title, cell.coordinate): cell.value for ws in workbook.worksheets for row in ws.iter_rows() for cell in row if isinstance(cell.value, str) and cell.value.startswith("=")}


def raw_formula_counts(package: zipfile.ZipFile) -> tuple[int, int, int]:
    records = 0
    groups: set[tuple[str, str]] = set()
    formulas = 0
    formula_tag = f"{{{MAIN_NS}}}f"
    for name in package.namelist():
        if not name.startswith("xl/worksheets/") or not name.endswith(".xml"):
            continue
        root = ET.fromstring(package.read(name))
        for formula in root.iter(formula_tag):
            formulas += 1
            if formula.get("t") == "shared":
                records += 1
                groups.add((name, formula.get("si", "")))
    return formulas, records, len(groups)


def normalized_worksheet(xml_bytes: bytes) -> bytes:
    root = ET.fromstring(xml_bytes)
    for formula in root.iter(f"{{{MAIN_NS}}}f"):
        formula.attrib.clear()
        formula.text = "__FORMULA_SERIALIZATION__"
    return ET.tostring(root, encoding="utf-8")


def _values_equal(left: object, right: object) -> bool:
    if isinstance(left, float) and isinstance(right, float) and math.isnan(left) and math.isnan(right):
        return True
    return left == right


def lineage_cells(workbook, sheet: str, ranges: tuple[str, ...]) -> list[str]:
    return [cell.coordinate for address in ranges for row in workbook[sheet][address] for cell in row]


def parser_independence_check() -> tuple[bool, str]:
    source = Path(__file__).read_text(encoding="utf-8")
    tree = ast.parse(source)
    imported = {
        node.module if isinstance(node, ast.ImportFrom) else alias.name
        for node in ast.walk(tree) if isinstance(node, (ast.Import, ast.ImportFrom))
        for alias in node.names
    }
    forbidden = sorted(name for name in imported if name.startswith("expand_shared_formulas") or name.startswith("lxml"))
    return (not forbidden, f"forbidden_shared_imports={forbidden}; validator=ElementTree+openpyxl")


def compare_native_excel_controls(reference: Path, recalculated: Path) -> dict:
    before = load_workbook(reference, data_only=True, read_only=False)["Controls"]
    after = load_workbook(recalculated, data_only=True, read_only=False)["Controls"]
    rows = [row for row in range(5, 76) if before[f"A{row}"].value not in (None, "") and before[f"C{row}"].value != "N/A"]
    differences = []
    for row in rows:
        for column in ("B", "C"):
            left, right = before[f"{column}{row}"].value, after[f"{column}{row}"].value
            if not _values_equal(left, right):
                differences.append({"cell": f"Controls!{column}{row}", "before": left, "after": right})
    return {"control_count": len(rows), "differences": differences, "passed": len(rows) == 66 and not differences}


def verify(before: Path, after: Path, desktop_recalc: Path | None = None) -> dict:
    failures: list[str] = []
    with zipfile.ZipFile(before) as left, zipfile.ZipFile(after) as right:
        left_names, right_names = left.namelist(), right.namelist()
        if left_names != right_names:
            failures.append("ZIP member order or membership changed")
        before_formula_count, before_records, before_groups = raw_formula_counts(left)
        after_formula_count, after_records, after_groups = raw_formula_counts(right)
        for name in left_names:
            left_data, right_data = left.read(name), right.read(name)
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                if normalized_worksheet(left_data) != normalized_worksheet(right_data):
                    failures.append(f"Non-formula worksheet XML changed: {name}")
            elif left_data != right_data:
                failures.append(f"Non-worksheet package member changed: {name} ({digest(left_data)} -> {digest(right_data)})")
    before_formulas = formula_map(before)
    after_formulas = formula_map(after)
    semantic_differences = sorted(f"{sheet}!{cell}" for sheet, cell in set(before_formulas) | set(after_formulas) if before_formulas.get((sheet, cell)) != after_formulas.get((sheet, cell)))
    before_values = workbook_cells(before, data_only=True)
    after_values = workbook_cells(after, data_only=True)
    cached_value_differences = sorted(f"{sheet}!{cell}" for sheet, cell in set(before_values) | set(after_values) if not _values_equal(before_values.get((sheet, cell)), after_values.get((sheet, cell))))
    left_wb = load_workbook(before, data_only=False, read_only=False)
    right_wb = load_workbook(after, data_only=False, read_only=False)
    right_cached = load_workbook(after, data_only=True, read_only=False)
    formula_errors = sorted(f"{ws.title}!{cell.coordinate}" for ws in right_cached.worksheets for row in ws.iter_rows() for cell in row if cell.value in ERRORS)
    missing_caches = sorted(f"{sheet}!{cell}" for sheet, cell in after_formulas if right_cached[sheet][cell].value is None)
    lineage = sorted(f"{sheet}!{cell}" for sheet, ranges in LINEAGE_RANGES.items() for cell in lineage_cells(right_wb, sheet, ranges))
    lineage_differences = sorted(cell for cell in lineage if cell in semantic_differences)
    parser_independent, parser_detail = parser_independence_check()
    native_controls = compare_native_excel_controls(after, desktop_recalc) if desktop_recalc else None
    checks = {
        "shared_formula_records_zero": after_records == 0,
        "shared_formula_groups_zero": after_groups == 0,
        "formula_count_unchanged": before_formula_count == after_formula_count == len(before_formulas) == len(after_formulas) == 8966,
        "resolved_formula_semantics_unchanged": not semantic_differences,
        "previously_affected_lineage_exact": len(lineage) == 25 and not lineage_differences,
        "cached_values_unchanged": not cached_value_differences,
        "tab_order_unchanged": left_wb.sheetnames == right_wb.sheetnames,
        "external_links_unchanged_and_zero": len(left_wb._external_links) == len(right_wb._external_links) == 0,
        "merged_cells_zero": sum(len(ws.merged_cells.ranges) for ws in right_wb.worksheets) == 0,
        "wrapped_populated_cells_zero": not any(cell.value is not None and cell.alignment.wrap_text for ws in right_wb.worksheets for row in ws.iter_rows() for cell in row),
        "formula_errors_zero": not formula_errors,
        "missing_formula_caches_zero": not missing_caches,
        "calculation_flags_unchanged": left_wb.calculation.__dict__ == right_wb.calculation.__dict__,
        "charts_layout_settings_unchanged": not failures,
    }
    if native_controls is not None:
        checks["native_excel_66_controls_unchanged"] = native_controls["passed"]
    return {
        "passed": all(checks.values()) and parser_independent, "checks": checks,
        "shared_formula_records_before": before_records, "shared_formula_records_after": after_records,
        "shared_formula_groups_before": before_groups, "shared_formula_groups_after": after_groups,
        "formula_count_before": before_formula_count, "formula_count_after": after_formula_count,
        "semantic_difference_count": len(semantic_differences), "semantic_differences": semantic_differences,
        "previously_affected_lineage_cell_count": len(lineage), "previously_affected_lineage_differences": lineage_differences,
        "cached_value_differences": cached_value_differences, "formula_errors": formula_errors,
        "missing_formula_caches": missing_caches, "package_failures": failures,
        "parser_independence": parser_detail, "native_excel_controls": native_controls,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("before", type=Path)
    parser.add_argument("after", type=Path)
    parser.add_argument("--desktop-recalc", type=Path)
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args()
    report = verify(args.before, args.after, args.desktop_recalc)
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(json.dumps(report, indent=2, sort_keys=True))
    return 0 if report["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
