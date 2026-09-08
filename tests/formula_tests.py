"""Formula integrity and pattern checks."""

from __future__ import annotations

import re
import zipfile
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from scripts.qa_common import finding, result


ERRORS = {"#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!"}


def _pattern(formula: str, row: int) -> str:
    return re.sub(r"([A-Z]{1,3})\$?\d+", lambda m: f"{m.group(1)}{{r{int(re.search(r'\d+', m.group(0)).group()) - row:+d}}}", formula)


def run(workbook: Path, manifest: dict, model_spec: dict) -> list[dict]:
    wb = load_workbook(workbook, data_only=False, read_only=False)
    cached = load_workbook(workbook, data_only=True, read_only=False)
    formulas = []
    missing_caches = []
    bad_colors = []
    wrapped = []
    dependencies = defaultdict(set)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if value is not None and cell.alignment.wrap_text:
                    wrapped.append(f"{ws.title}!{cell.coordinate}")
                if isinstance(value, str) and value.startswith("="):
                    formulas.append((ws.title, cell.coordinate, value))
                    if cached[ws.title][cell.coordinate].value is None:
                        missing_caches.append(f"{ws.title}!{cell.coordinate}")
                    color = cell.font.color
                    rgb = color.rgb if color and color.type == "rgb" else None
                    expected_color = model_spec.get("formula_colors", {}).get("cross_sheet_formula" if "!" in value else "local_formula")
                    if rgb != expected_color:
                        bad_colors.append(f"{ws.title}!{cell.coordinate}")
                    for ref in re.findall(r"(?:'([^']+)'|([A-Za-z0-9_ ]+))!\$?[A-Z]{1,3}\$?\d+", value):
                        dependencies[ws.title].add(ref[0] or ref[1])
    with zipfile.ZipFile(workbook) as package:
        shared = sum(package.read(name).count(b't="shared"') for name in package.namelist() if name.startswith("xl/worksheets/") and name.endswith(".xml"))
    long_formulas = [f"{s}!{c}" for s, c, value in formulas if len(value) > 1000]
    error_cells = [f"{ws.title}!{cell.coordinate}" for ws in wb.worksheets for row in ws.iter_rows() for cell in row if cell.value in ERRORS]
    patterns = sorted({f"{sheet}:{_pattern(value, int(re.search(r'\d+', coord).group()))}" for sheet, coord, value in formulas})
    expected_patterns = manifest.get("formula_patterns", patterns)
    checks = [
        result("shared_formulas_zero", shared == 0, f"shared_formulas={shared}"),
        result("formula_errors_zero", not error_cells, f"errors={error_cells}"),
        result("missing_caches_zero", not missing_caches, f"missing_caches={missing_caches}"),
        result("long_formulas_zero", not long_formulas, f"long_formulas={long_formulas}"),
        result("wrapped_populated_cells_zero", not wrapped, f"wrapped_cells={wrapped}"),
        result("formula_color_convention", not bad_colors, f"bad_formula_colors={bad_colors}"),
        result("dependency_graph", all(dep in wb.sheetnames for deps in dependencies.values() for dep in deps), f"dependencies={dict((k, sorted(v)) for k, v in dependencies.items())}"),
        result("formula_pattern_regression", patterns == expected_patterns, "Formula patterns must match manifest baseline"),
    ]
    for index, item in enumerate(checks, 1):
        if not item["passed"]:
            item["findings"] = [finding(f"FRM-{index:03d}", item["test"], item["details"], "Correct formulas or formatting", item["test"])]
    return checks
