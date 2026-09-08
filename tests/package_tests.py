"""Workbook package, presentation, and manifest checks."""

from __future__ import annotations

import zipfile
from pathlib import Path

from openpyxl import load_workbook

from scripts.qa_common import file_sha256, finding, result


def run(workbook: Path, manifest: dict, model_spec: dict) -> list[dict]:
    wb = load_workbook(workbook, data_only=False, read_only=False)
    expected_tabs = model_spec.get("tab_order", [])
    expected_hash = manifest.get("workbook_sha256")
    approved_hash = model_spec.get("base_workbook", {}).get("sha256")
    results = [
        result("exact_tab_order", bool(expected_tabs) and wb.sheetnames == expected_tabs, f"actual={wb.sheetnames}; expected={expected_tabs}"),
        result("external_links_zero", len(wb._external_links) == 0, f"external_links={len(wb._external_links)}"),
        result("merged_cells_zero", sum(len(ws.merged_cells.ranges) for ws in wb.worksheets) == 0, "Merged ranges must be zero"),
        result("manifest_reconciliation", manifest.get("workbook") == workbook.name, "Manifest workbook name must match candidate"),
        result("version_hash_checks", bool(expected_hash) and expected_hash == approved_hash == file_sha256(workbook), "Workbook, manifest, and approved Base SHA-256 must match"),
    ]
    with zipfile.ZipFile(workbook) as package:
        names = set(package.namelist())
        chart_parts = [name for name in names if name.startswith("xl/charts/")]
        drawing_refs = [name for name in names if name.startswith("xl/drawings/")]
        results.append(result("chart_package_checks", not chart_parts or bool(drawing_refs), f"charts={len(chart_parts)} drawings={len(drawing_refs)}"))
    for index, item in enumerate(results, 1):
        if not item["passed"]:
            item["findings"] = [finding(f"PKG-{index:03d}", item["test"], item["details"], "Rebuild candidate package or manifest", item["test"])]
    return results
